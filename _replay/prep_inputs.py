"""Rebuild all replay inputs into the permanent _replay/ folder.
- yard_2324.json : transcribed from DocScanner PDF (full season, 138 days) + cross-check vs recovered
- yard_2425.json : from DAILY INDENT PAD REC. - 2024-25.xlsx (full season)
- yard_2526.json : from YARD POSITION FOR CRUSHING 2025-26.xls (full season)
- per-centre CSVs for 23-24, 24-25 (2023-24-25.xlsx) and 25-26 (Indent_Purchase_DataPasteOnly_2025.xlsx)
"""
import openpyxl, xlrd, csv, os, json, re
from datetime import datetime
from collections import defaultdict

ROOT = r'c:/Users/tejas.rai/Desktop/Projects/GannaApp'
REP  = os.path.join(ROOT, '_replay')
SD   = os.path.join(REP, 'season_data'); os.makedirs(SD, exist_ok=True)

def dd_mm_yy(s):
    d, m, y = s.split('.')
    return f'20{y}-{m.zfill(2)}-{d.zfill(2)}'

# ─────────────── 2023-24 yard: transcribed from the PDF (Date, CaneReq, YardGate8AM, YardCentre8AM) ───────────────
PDF_2324 = [
 ("13.11.23",0,0,0),("14.11.23",40000,2285,0),("15.11.23",50000,0,0),("16.11.23",80000,29015,25555),
 ("17.11.23",50000,31093,18781),("18.11.23",50000,31856,20707),("19.11.23",50000,32953,31253),("20.11.23",50000,26189,33567),
 ("21.11.23",50000,24494,31075),("22.11.23",50000,26832,33968),("23.11.23",65000,15528,20252),("24.11.23",75000,10390,10356),
 ("25.11.23",85000,11904,11693),("26.11.23",90000,25709,24561),("27.11.23",80000,21477,12295),("28.11.23",80000,25183,10771),
 ("29.11.23",80000,27173,14285),("30.11.23",80000,27283,16696),("01.12.23",85000,20638,19460),("02.12.23",0,30019,40137),
 ("03.12.23",60000,33663,41502),("04.12.23",90000,25042,18516),("05.12.23",0,18339,16677),("06.12.23",90000,9694,14595),
 ("07.12.23",96000,5249,12424),("08.12.23",96000,3759,4937),("09.12.23",96000,8619,6842),("10.12.23",96000,15886,15521),
 ("11.12.23",96000,22224,18436),("12.12.23",96000,25340,16058),("13.12.23",96000,27363,14447),("14.12.23",96000,19993,12901),
 ("15.12.23",96000,17104,7079),("16.12.23",96000,17053,5232),("17.12.23",96000,13654,1983),("18.12.23",96000,10499,4268),
 ("19.12.23",96000,12230,11422),("20.12.23",96000,16729,13651),("21.12.23",96000,16285,15709),("22.12.23",96000,16359,16606),
 ("23.12.23",96000,18431,9649),("24.12.23",96000,12257,1267),("25.12.23",45000,12658,13175),("26.12.23",96000,12820,4777),
 ("27.12.23",96000,8771,9501),("28.12.23",96000,1550,8055),("29.12.23",96000,218,0),("30.12.23",96000,4000,6251),
 ("31.12.23",96000,9287,3002),("01.01.24",96000,9963,1468),("02.01.24",96000,15920,5297),("03.01.24",96000,24139,8703),
 ("04.01.24",96000,31401,19526),("05.01.24",96000,34225,14548),("06.01.24",96000,29476,15636),("07.01.24",96000,24626,14034),
 ("08.01.24",96000,13232,9692),("09.01.24",96000,8140,3107),("10.01.24",96000,10416,7706),("11.01.24",96000,7537,6000),
 ("12.01.24",96000,8995,8115),("13.01.24",96000,12569,7355),("14.01.24",96000,5481,6330),("15.01.24",96000,6201,7376),
 ("16.01.24",96000,12465,8871),("17.01.24",30000,19278,30932),("18.01.24",50000,29516,32891),("19.01.24",80000,32008,21749),
 ("20.01.24",96000,25414,19701),
 ("21.01.24",96000,13858,22398),("22.01.24",96000,11694,17836),("23.01.24",96000,16855,23346),("24.01.24",96000,21172,15259),
 ("25.01.24",96000,25665,22219),("26.01.24",96000,22193,14791),("27.01.24",96000,13540,21403),("28.01.24",96000,14265,14126),
 ("29.01.24",96000,9911,10655),("30.01.24",96000,12708,10977),("31.01.24",96000,11779,6263),("01.02.24",96000,6899,8007),
 ("02.02.24",96000,9831,7628),("03.02.24",96000,10857,10969),("04.02.24",96000,2290,7162),("05.02.24",96000,6797,14581),
 ("06.02.24",96000,8277,17716),("07.02.24",96000,8283,19235),("08.02.24",96000,15858,35923),("09.02.24",96000,11362,20617),
 ("10.02.24",96000,7342,21520),("11.02.24",96000,4917,20196),("12.02.24",96000,5701,14463),("13.02.24",96000,16295,28867),
 ("14.02.24",96000,11452,18338),("15.02.24",96000,7259,11635),("16.02.24",96000,2554,3698),("17.02.24",96000,6448,2584),
 ("18.02.24",96000,10818,5991),("19.02.24",96000,15689,4012),("20.02.24",96000,23301,10914),("21.02.24",96000,28939,16890),
 ("22.02.24",96000,29748,21965),("23.02.24",96000,27505,21301),("24.02.24",96000,24667,18212),("25.02.24",96000,16980,18132),
 ("26.02.24",96000,6024,10360),("27.02.24",96000,2558,1029),("28.02.24",96000,2428,3687),("29.02.24",96000,1097,4115),
 ("01.03.24",96000,10670,25233),("02.03.24",96000,559,7075),("03.03.24",96000,2730,507),("04.03.24",96000,15091,6935),
 ("05.03.24",96000,21108,8839),("06.03.24",96000,23915,9533),("07.03.24",96000,23875,8925),("08.03.24",96000,22412,10128),
 ("09.03.24",96000,17210,11777),("10.03.24",96000,15950,9191),("11.03.24",96000,5014,3095),("12.03.24",96000,4381,3825),
 ("13.03.24",96000,187,3945),("14.03.24",96000,1009,6959),("15.03.24",96000,5371,8704),("16.03.24",96000,8934,10900),
 ("17.03.24",96000,13613,17903),("18.03.24",96000,21096,15080),("19.03.24",96000,22010,15517),("20.03.24",96000,21266,17119),
 ("21.03.24",96000,24275,23129),("22.03.24",96000,27859,28255),("23.03.24",96000,30945,28812),("24.03.24",96000,17842,21860),
 ("25.03.24",96000,3398,529),("26.03.24",96000,6803,10772),("27.03.24",96000,7256,23027),("28.03.24",96000,9288,20396),
 ("29.03.24",96000,16764,28688),
]
yard_2324 = {}
for s, req, yg, yc in PDF_2324:
    yard_2324[dd_mm_yy(s)] = {'req': float(req), 'yardGate': float(yg), 'yardCentre': float(yc)}
json.dump(yard_2324, open(os.path.join(REP, 'yard_2324.json'), 'w'), indent=1)
print(f'yard_2324 (PDF): {len(yard_2324)} days, {min(yard_2324)}..{max(yard_2324)}')

# cross-check overlap vs the recovered-from-Excel yard (decision-day keyed)
rec = json.load(open(os.path.join(REP, 'yard_202324_recovered.json')))
mism = 0; checked = 0
for k, v in rec.items():
    if k in yard_2324 and v.get('yardGate') is not None:
        checked += 1
        if abs(yard_2324[k]['yardGate'] - v['yardGate']) > 1 or abs(yard_2324[k]['yardCentre'] - (v['yardCentre'] or 0)) > 1:
            mism += 1
            if mism <= 5: print(f'  MISMATCH {k}: PDF gate={yard_2324[k]["yardGate"]} centre={yard_2324[k]["yardCentre"]} vs recovered gate={v["yardGate"]} centre={v["yardCentre"]}')
print(f'cross-check vs recovered: {checked} overlapping days, {mism} mismatches')

# ─────────────── 2024-25 yard from DAILY INDENT PAD REC. xlsx ───────────────
def parse_dmy_any(s):
    m = re.match(r'^(\d{1,2})\.+(\d{1,2})\.+(\d{2,4})$', str(s).strip())
    if not m: return None
    d, mo, y = m.groups(); y = y[-2:]
    return f'20{y}-{mo.zfill(2)}-{d.zfill(2)}'
wb = openpyxl.load_workbook(os.path.join(ROOT, 'DAILY INDENT PAD REC. - 2024-25.xlsx'), read_only=True, data_only=True)
ws = wb['date wise indent 24.25']; yard_2425 = {}
for row in ws.iter_rows(min_row=5, values_only=True):
    iso = parse_dmy_any(row[1]) if row[1] is not None else None
    if not iso: continue
    def num(v):
        try: return float(v)
        except: return None
    yard_2425[iso] = {'req': num(row[2]), 'yardGate': num(row[9]), 'yardCentre': num(row[10])}
wb.close()
json.dump(yard_2425, open(os.path.join(REP, 'yard_2425.json'), 'w'), indent=1)
print(f'yard_2425 (xlsx): {len(yard_2425)} days, {min(yard_2425)}..{max(yard_2425)}')

# ─────────────── 2025-26 yard from YARD POSITION xls ───────────────
wb = xlrd.open_workbook(os.path.join(ROOT, 'YARD POSITION FOR CRUSHING 2025-26.xls'))
sh = wb.sheet_by_name('Indent Crushing Bal.Rep-025-026'); yard_2526 = {}
for r in range(5, sh.nrows):
    iso = parse_dmy_any(sh.cell_value(r, 1))
    if not iso: continue
    def num(c):
        try: return float(sh.cell_value(r, c))
        except: return None
    yard_2526[iso] = {'req': num(2), 'yardGate': num(9), 'yardCentre': num(10)}
json.dump(yard_2526, open(os.path.join(REP, 'yard_2526.json'), 'w'), indent=1)
print(f'yard_2526 (xls): {len(yard_2526)} days, {min(yard_2526)}..{max(yard_2526)}')

# ─────────────── per-centre CSVs ───────────────
def to_date(v):
    if isinstance(v, datetime): return v.date()
    if v is None: return None
    s = str(v).strip()
    for f in ('%Y-%m-%d', '%d-%m-%Y'):
        try: return datetime.strptime(s[:10], f).date()
        except: pass
    return None
def season_of(d):
    y = d.year
    return f'{y%100:02d}-{(y+1)%100:02d}' if d.month >= 9 else f'{(y-1)%100:02d}-{y%100:02d}'
def dmy(d): return d.strftime('%d-%m-%Y')

def export_from(wbpath, season_filter, out_prefix):
    wb = openpyxl.load_workbook(wbpath, read_only=True, data_only=True)
    names = {}; si = []; sp = []
    for row in wb['Indent'].iter_rows(min_row=2, values_only=True):
        if row[0] is None or str(row[0]).strip() == '': continue
        d = to_date(row[2]); q = row[4]
        if d is None or not isinstance(q, (int, float)): continue
        if season_filter and season_of(d) != season_filter: continue
        c = str(row[0]).strip(); n = str(row[1]).strip() if row[1] else ''
        if n: names.setdefault(c, n)
        si.append((c, n, d, float(q)))
    for row in wb['Purchase'].iter_rows(min_row=2, values_only=True):
        if row[0] is None or str(row[0]).strip() == '': continue
        pd = to_date(row[2]); idd = to_date(row[3]); q = row[5]
        if pd is None or not isinstance(q, (int, float)): continue
        if season_filter and season_of(pd) != season_filter: continue
        sp.append((str(row[0]).strip(), pd, idd, float(q)))
    master = {}
    if 'Center Data' in wb.sheetnames:
        for row in wb['Center Data'].iter_rows(min_row=2, values_only=True):
            if row[0] is None: continue
            b = row[10] if len(row) > 10 else None
            try: b = float(b)
            except: b = None
            master[str(row[0]).strip()] = b
    wb.close()
    with open(os.path.join(SD, f'{out_prefix}_indent.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['Code', 'Center Name', 'Indent Date', 'No of Purchy', 'Qty in Qtls'])
        for c, n, d, q in si: w.writerow([c, n or names.get(c, c), dmy(d), '', q])
    with open(os.path.join(SD, f'{out_prefix}_purchase.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['Code', 'Center Name', 'Purchase Date', 'Indent Date', 'No of Purchy', 'Qty in Qtls'])
        for c, pd, idd, q in sp: w.writerow([c, names.get(c, c), dmy(pd), dmy(idd) if idd else '', '', q])
    tot = defaultdict(float)
    for c, n, d, q in si: tot[c] += q
    with open(os.path.join(SD, f'{out_prefix}_bonding.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['Code', 'Center', 'Bonding'])
        for c in sorted(tot, key=lambda x: -tot[x]):
            b = master.get(c)
            if not b or b <= 0: b = tot[c]
            w.writerow([c, names.get(c, c), b])
    print(f'  {out_prefix}: indents={len(si):,} purchases={len(sp):,} centres={len(tot)}')

print('exporting per-centre CSVs...')
export_from(os.path.join(ROOT, '2023-24-25.xlsx'), '23-24', '23-24')
export_from(os.path.join(ROOT, '2023-24-25.xlsx'), '24-25', '24-25')
export_from(os.path.join(ROOT, 'Indent_Purchase_DataPasteOnly_2025.xlsx'), None, '25-26')
print('DONE — all inputs in _replay/')
