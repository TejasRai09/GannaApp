import openpyxl, sys
from datetime import datetime, timedelta
from collections import defaultdict
sys.stdout.reconfigure(encoding='utf-8')

XLSM = "IndentingModel_TestFilexlsm.xlsm"
T = datetime(2023, 1, 27)

wb = openpyxl.load_workbook(XLSM, read_only=True, data_only=True)

# Check Centers sheet
ws_c = wb['Centers']
print("=== Centers sheet (first 30 rows) ===")
for i, row in enumerate(ws_c.iter_rows(min_col=1, max_col=10, values_only=True)):
    if i >= 30: break
    if any(v is not None for v in row):
        print(f"  Row {i+1}: {list(row)}")

# Check total indents for ALL centers on key dates
ws_i = wb['Indent']
all_centers = defaultdict(lambda: defaultdict(float))  # [date][code] = qty
for row in ws_i.iter_rows(min_row=2, values_only=True):
    code = row[0]; d = row[2]; qty = row[4]
    if code is None or d is None: continue
    if isinstance(d, datetime) and datetime(2023,1,20) <= d <= datetime(2023,1,30):
        all_centers[d.date()][int(code)] = float(qty or 0)

print("\n=== All centers with indents on Jan 20, 27, 28, 29 ===")
for date in [datetime(2023,1,20).date(), datetime(2023,1,27).date(),
             datetime(2023,1,28).date(), datetime(2023,1,29).date()]:
    total = sum(all_centers[date].values())
    gate_codes = [(c, v) for c, v in all_centers[date].items()
                  if c in [1, 26, 27] or (v > 1000)]
    print(f"  {date}: total={total:,.0f}, GATE-area codes: {gate_codes[:5]}")

# Also check: what is the SUM of ALL center indents for Jan 27?
print(f"\n  Jan 27 all codes: {dict(all_centers[datetime(2023,1,27).date()])}")

# Check what "GATE" in the forecast = which codes?
# Cross-check: Forecast says Jan 27 arrival = 32,517
# Let's see which combination of codes sums to 32,517
jan27 = all_centers[datetime(2023,1,27).date()]
code1 = jan27.get(1, 0)
code26 = jan27.get(26, 0)
code27 = jan27.get(27, 0)
print(f"\n  Code 1 (GATE): {code1}")
print(f"  Code 26 (GATE LMP): {code26}")
print(f"  Code 27 (GATE ARNI): {code27}")
print(f"  Sum 1+26+27: {code1+code26+code27}")
print(f"  Forecast expects: 32,517")
