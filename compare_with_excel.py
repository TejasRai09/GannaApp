"""
compare_with_excel.py
Compares GannaApp algorithm output vs IndentingModel_TestFilexlsm.xlsm
Test case: January 27, 2023 (Day 1 Report)
Run: python compare_with_excel.py
"""

import openpyxl
from datetime import datetime, timedelta, timezone
from collections import defaultdict

XLSM = "IndentingModel_TestFilexlsm.xlsm"

# ─── Parameters (from Report sheet, Row 8-13) ────────────────────────────────
T             = datetime(2023, 1, 27)
PLANT_START   = datetime(2022, 11, 15)
DAILY_REQ     = 100_000          # max daily capacity
PLANT_CAP     = 0.80             # plant capacity % → effective = 80,000
STD_GATE      = 10_000
AVAIL_GATE    = 4_000
STD_CENTRE    = 6_000
AVAIL_CENTRE  = 8_000
# ─────────────────────────────────────────────────────────────────────────────

def to_date(v):
    if isinstance(v, datetime):
        return v.replace(hour=0, minute=0, second=0, microsecond=0)
    return None

print("Loading Excel workbook …")
wb = openpyxl.load_workbook(XLSM, read_only=True, data_only=True)

# ── 1. Read Bonding ──────────────────────────────────────────────────────────
bonding = {}   # code → {name, qty, is_gate}
ws_b = wb['Bonding']
for row in ws_b.iter_rows(min_row=2, values_only=True):
    code = row[0]
    name = row[1]
    qty  = row[2]
    if code is None or not isinstance(code, (int, float)):
        continue
    bonding[int(code)] = {
        'name': str(name).strip(),
        'qty':  float(qty or 0),
        'is_gate': 'GATE' in str(name).upper()
    }
print(f"  Bonding: {len(bonding)} centers")

total_bonding     = sum(v['qty'] for v in bonding.values())
total_gate_bond   = sum(v['qty'] for v in bonding.values() if     v['is_gate'])
total_centre_bond = sum(v['qty'] for v in bonding.values() if not v['is_gate'])

# ── 2. Read Indent data ──────────────────────────────────────────────────────
# indent[(code, date)] = qty
indents = defaultdict(float)
ws_i = wb['Indent']
for row in ws_i.iter_rows(min_row=2, values_only=True):
    code = row[0]; d = to_date(row[2]); qty = row[4]
    if code is None or d is None:
        continue
    indents[(int(code), d)] += float(qty or 0)
print(f"  Indents: {len(indents)} records")

# ── 3. Read Purchase data ────────────────────────────────────────────────────
# purchases[(code, indent_date, purchase_date)] = qty
purchases_raw = []
ws_p = wb['Purchase']
for row in ws_p.iter_rows(min_row=2, values_only=True):
    code = row[0]; pur_d = to_date(row[2]); ind_d = to_date(row[3]); qty = row[5]
    if code is None or pur_d is None or ind_d is None:
        continue
    purchases_raw.append((int(code), ind_d, pur_d, float(qty or 0)))
print(f"  Purchases: {len(purchases_raw)} records")

# ── 4. Expected results from Report sheet ────────────────────────────────────
expected = {}  # code → indent_automatic
ws_r = wb['Report']
# Columns J-L (0-indexed 9-11): Centre Code, Center Name, Indent Automatic
for row in ws_r.iter_rows(min_row=5, values_only=True):
    code = row[9]; name = row[10]; val = row[11]
    if code is None or not isinstance(code, (int, float)):
        continue
    expected[int(code)] = float(val or 0)
print(f"  Expected results: {len(expected)} centers")
wb.close()

# ─────────────────────────────────────────────────────────────────────────────
# ALGORITHM (mirrors calculationService.ts exactly)
# ─────────────────────────────────────────────────────────────────────────────

effective_req = DAILY_REQ * PLANT_CAP   # 80,000

# ── Overrun (season-wide) ────────────────────────────────────────────────────
total_ind_qty = sum(
    qty for (code, ind_d), qty in indents.items()
    if ind_d >= PLANT_START and ind_d <= T
)
total_pur_qty = sum(
    qty for (code, ind_d, pur_d, qty) in purchases_raw
    if ind_d >= PLANT_START and pur_d <= T
)
overrun = (total_pur_qty / total_ind_qty - 1) if total_ind_qty > 0 else 0
print(f"\nSeason overrun: {overrun:.4f} ({overrun*100:.2f}%)")
print(f"Total indents:  {total_ind_qty:,.0f} Qtl")
print(f"Total purchases:{total_pur_qty:,.0f} Qtl")

# ── D-weights: closed indents T-7 to T-4 ────────────────────────────────────
def compute_d_weights(code, t, indents, purchases_raw):
    """Compute avg D1-D4 weights from last 4 closed indents (T-7 to T-4)."""
    t4 = t - timedelta(days=4)
    t7 = t - timedelta(days=7)

    # Collect per-indent D rates in date order
    d_rates = []
    for delta in range(4):  # T-4, T-5, T-6, T-7
        ind_date = t - timedelta(days=4+delta)
        if ind_date < t7:
            break
        ind_qty = indents.get((code, ind_date), 0)
        if ind_qty <= 0:
            continue
        # Bucket purchases
        d1 = d2 = d3 = d4 = 0.0
        for (pc, pid, ppd, pqty) in purchases_raw:
            if pc != code or pid != ind_date:
                continue
            diff = (ppd - ind_date).days
            if diff <= 0:  d1 += pqty
            elif diff == 1: d2 += pqty
            elif diff == 2: d3 += pqty
            else:           d4 += pqty
        d_rates.append({
            'd1': d1 / ind_qty,
            'd2': d2 / ind_qty,
            'd3': d3 / ind_qty,
            'd4': d4 / ind_qty,
        })

    if not d_rates:
        return None  # no recent data

    n = len(d_rates)
    return {
        'd1': sum(r['d1'] for r in d_rates) / n,
        'd2': sum(r['d2'] for r in d_rates) / n,
        'd3': sum(r['d3'] for r in d_rates) / n,
        'd4': sum(r['d4'] for r in d_rates) / n,
    }

def compute_season_d_weights(code, plant_start, t, indents, purchases_raw):
    """Season-wide D-weight fallback."""
    d_rates = []
    t4 = t - timedelta(days=4)
    for (c, ind_d), ind_qty in indents.items():
        if c != code or ind_d < plant_start or ind_d > t4 or ind_qty <= 0:
            continue
        d1 = d2 = d3 = d4 = 0.0
        for (pc, pid, ppd, pqty) in purchases_raw:
            if pc != code or pid != ind_d:
                continue
            diff = (ppd - ind_d).days
            if diff <= 0:  d1 += pqty
            elif diff == 1: d2 += pqty
            elif diff == 2: d3 += pqty
            else:           d4 += pqty
        if ind_qty > 0:
            d_rates.append({'d1': d1/ind_qty, 'd2': d2/ind_qty,
                            'd3': d3/ind_qty, 'd4': d4/ind_qty})
    if not d_rates:
        return {'d1': 0, 'd2': 0, 'd3': 0, 'd4': 0}
    n = len(d_rates)
    return {
        'd1': sum(r['d1'] for r in d_rates) / n,
        'd2': sum(r['d2'] for r in d_rates) / n,
        'd3': sum(r['d3'] for r in d_rates) / n,
        'd4': sum(r['d4'] for r in d_rates) / n,
    }

# ── Run calculation per center ───────────────────────────────────────────────
T0 = T
T1 = T + timedelta(days=1)
T2 = T + timedelta(days=2)

results = []
for code, b in bonding.items():
    if b['qty'] <= 0:
        continue

    # D-weights
    w = compute_d_weights(code, T0, indents, purchases_raw)
    if w is None or w['d1'] < 0.05:
        w = compute_season_d_weights(code, PLANT_START, T0, indents, purchases_raw)

    # Requirement by bonding
    req_by_bonding = effective_req * (b['qty'] / total_bonding)

    # Stock adjustment
    if b['is_gate']:
        stock_diff = STD_GATE - AVAIL_GATE
        stock_adj  = stock_diff * (b['qty'] / total_gate_bond)
    else:
        stock_diff = STD_CENTRE - AVAIL_CENTRE
        stock_adj  = stock_diff * (b['qty'] / total_centre_bond)

    adj_req = req_by_bonding + stock_adj

    # Forecast T+3 from open indents
    i_t0 = indents.get((code, T0), 0)
    i_t1 = indents.get((code, T1), 0)
    i_t2 = indents.get((code, T2), 0)
    forecast_t3 = i_t2 * w['d2'] + i_t1 * w['d3'] + i_t0 * w['d4']

    # Final indent
    gap = max(0, adj_req - forecast_t3)
    target_arrival = gap / (1 + overrun) if overrun > -1 else 0
    final_indent = target_arrival / w['d1'] if w['d1'] > 0 else 0

    results.append({
        'code': code,
        'name': b['name'],
        'app_indent': final_indent,
        'expected':   expected.get(code, None),
        'adj_req':    adj_req,
        'forecast':   forecast_t3,
        'gap':        gap,
        'd1':         w['d1'],
    })

# ── Print comparison ─────────────────────────────────────────────────────────
results.sort(key=lambda x: x['code'])

matched  = 0
close    = 0
mismatch = 0
no_exp   = 0

print(f"\n{'='*85}")
print(f"{'Code':>5} {'Center':<30} {'App Indent':>12} {'Excel Exp':>12} {'Diff':>8} {'Status'}")
print(f"{'='*85}")

for r in results:
    app  = r['app_indent']
    exp  = r['expected']

    if exp is None:
        status = 'NO EXP'
        no_exp += 1
        diff_s = '-'
    else:
        diff = abs(app - exp)
        pct  = (app - exp) / exp * 100 if exp != 0 else float('inf')
        diff_s = f"{pct:+.1f}%"
        if diff < 5:
            status = 'MATCH  OK'
            matched += 1
        elif diff < 100:
            status = 'CLOSE  ~'
            close += 1
        else:
            status = 'MISMATCH X'
            mismatch += 1

    exp_s = f"{exp:>12.1f}" if exp is not None else f"{'N/A':>12}"
    print(f"{r['code']:>5} {r['name']:<30} {app:>12.1f} {exp_s} {diff_s:>8}  {status}")

print(f"{'='*85}")
print(f"\nSUMMARY: MATCH={matched}  CLOSE={close}  MISMATCH={mismatch}  NO_EXPECTED={no_exp}")
print(f"Total centers processed: {len(results)}")

# ── GATE detailed breakdown ──────────────────────────────────────────────────
gate = next((r for r in results if r['code'] == 1), None)
if gate:
    print(f"\n--- GATE Detailed Breakdown ---")
    print(f"  D1 weight:        {gate['d1']:.5f}")
    print(f"  Adj requirement:  {gate['adj_req']:,.2f}")
    print(f"  Forecast T+3:     {gate['forecast']:,.2f}")
    print(f"  Gap:              {gate['gap']:,.2f}")
    print(f"  Overrun:          {overrun*100:.4f}%")
    print(f"  App final indent: {gate['app_indent']:,.2f}")
    print(f"  Excel expected:   {gate['expected']:,.2f}")
    print(f"  Difference:       {abs(gate['app_indent'] - gate['expected']):.2f} Qtl")
