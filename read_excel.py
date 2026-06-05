import openpyxl, warnings, sys, io
warnings.filterwarnings('ignore')
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def dump_file(path, max_rows=80, max_cols=25):
    wb = openpyxl.load_workbook(path, data_only=True, keep_vba=False)
    print(f'\n===== FILE: {path} =====')
    print(f'Sheets: {wb.sheetnames}')
    for name in wb.sheetnames:
        ws = wb[name]
        print(f'\n--- Sheet: {name} (dims: {ws.dimensions}) ---')
        for r in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
            if any(c is not None for c in r):
                print(list(r))

base = r'c:\Users\tejas.rai\Desktop\Projects\GannaApp\Current Season Files for Testing'
for f in sys.argv[1:]:
    dump_file(f'{base}\\{f}')
